"""Build a restrained scientific Word manuscript from the review Markdown."""
import argparse
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def build(source, package, target):
    d=Document()
    sec=d.sections[0]
    sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
    sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
    sec.header_distance=Inches(.26);sec.footer_distance=Inches(.28)
    for name in ('Normal','Title','Heading 1','Heading 2','Caption'):
        style=d.styles[name]
        style.font.name='Liberation Serif'
        style.font.color.rgb=RGBColor(0,0,0)
        style.font.underline=False
        # The bundled Word template has a blue Title border and theme fonts.
        for border in list(style.element.xpath('./w:pPr/w:pBdr')):
            border.getparent().remove(border)
        fonts=style.element.find(qn('w:rPr')).find(qn('w:rFonts'))
        if fonts is not None:
            for key in list(fonts.attrib):
                if 'theme' in key.lower(): del fonts.attrib[key]
    n=d.styles['Normal'];n.font.size=Pt(11)
    n.paragraph_format.line_spacing=1.06;n.paragraph_format.space_after=Pt(6)
    n.paragraph_format.widow_control=True
    t=d.styles['Title'];t.font.size=Pt(20);t.font.bold=True;t.font.name='Liberation Sans'
    t.paragraph_format.space_after=Pt(8)
    for name,size in (('Heading 1',14),('Heading 2',11.6)):
        h=d.styles[name];h.font.size=Pt(size);h.font.bold=True;h.font.name='Liberation Sans'
        h.paragraph_format.space_before=Pt(9);h.paragraph_format.space_after=Pt(5)
        h.paragraph_format.keep_with_next=True
    cap=d.styles['Caption'];cap.font.size=Pt(9.2);cap.font.italic=False;cap.font.bold=False
    cap.paragraph_format.space_after=Pt(7)
    header=sec.header.paragraphs[0];header.text='FAIRLENS REPLICATION   •   AUTHOR REVIEW'
    header.runs[0].font.name='Liberation Sans';header.runs[0].font.size=Pt(8)
    header.runs[0].font.color.rgb=RGBColor.from_string('555555')
    footer=sec.footer.paragraphs[0];footer.alignment=2
    r=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
    d.core_properties.title='Refusals and language sensitivity in a FairLens replication with typed decision models'
    d.core_properties.subject='Complete closed question replication with English and Brazilian Portuguese prompts'
    d.core_properties.author=''
    lines=source.read_text().splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1;continue
        if line=='<!-- pagebreak -->':d.add_page_break();i+=1;continue
        if line.startswith('# '):d.add_paragraph(line[2:],style='Title');i+=1;continue
        if line.startswith('## '):d.add_paragraph(line[3:],style='Heading 1');i+=1;continue
        if line.startswith('### '):d.add_paragraph(line[4:],style='Heading 2');i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',c) for c in cells): rows.append(cells)
                i+=1
            table=d.add_table(rows=0,cols=len(rows[0]));table.autofit=False
            widths=([1.25,1.66,1.72,1.42,.78] if len(rows[0])==5 else [1.22,1.02,1.02,1.16,1.19,1.22])
            for col,width in zip(table.columns,widths):col.width=Inches(width)
            pr=table._tbl.tblPr
            borders=OxmlElement('w:tblBorders')
            for edge in ('top','left','bottom','right','insideH','insideV'):
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D6D6D6');borders.append(el)
            pr.append(borders)
            for j,values in enumerate(rows):
                row=table.add_row()
                cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
                if j==0:
                    repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
                for k,(cell,value) in enumerate(zip(row.cells,values)):
                    cell.width=Inches(widths[k]);p=cell.paragraphs[0]
                    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
                    p.paragraph_format.line_spacing=1.02
                    if k:p.alignment=2
                    run=p.add_run(value);run.font.size=Pt(9.4);run.bold=j==0
                    if j==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F1F1F1');cell._tc.get_or_add_tcPr().append(shade)
            continue
        image=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line)
        if image:
            p=d.add_paragraph();p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(package/image.group(2)),width=Inches(6.7))
            i+=1;continue
        p=d.add_paragraph(line,style='Caption' if re.match(r'^(Table|Figure) \d+\.',line) else None)
        if line.startswith('Scientific manuscript'):
            for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('555555')
        if line.startswith('['):
            p.paragraph_format.space_after=Pt(6)
            for r in p.runs:r.font.size=Pt(9.5)
        i+=1
    target.parent.mkdir(parents=True,exist_ok=True)
    d.save(target)
    print(target)


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--source',type=Path,required=True);p.add_argument('--package',type=Path,required=True);p.add_argument('--out',type=Path,required=True)
    a=p.parse_args();build(a.source,a.package,a.out)
